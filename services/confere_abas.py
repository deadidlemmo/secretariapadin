"""Leitura das Listas Piloto que possuem uma aba visivel para cada turma."""

import re
import unicodedata
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


_HEADER_ALIASES = {
    "numero": {"N", "NO", "NUMERO"},
    "rm": {"RM"},
    "nome": {"NOME", "NOMEDOALUNO", "ALUNO"},
    "data_nascimento": {"NASC", "NASCIMENTO", "DATANASC", "DATADENASCIMENTO"},
    "ra": {"RA"},
    "digito_ra": {"DIG", "DIGITO", "DIGITORA", "DIGITODORA"},
    "situacao": {"COD", "CODIGO", "SITUACAO", "STATUS"},
    "observacoes": {"OBS", "OBSERVACAO", "OBSERVACOES"},
}
_REQUIRED_HEADERS = {"numero", "rm", "nome", "data_nascimento", "ra", "situacao"}
_INVALID_NAMES = {
    "0", "NAN", "NONE", "NAT", "NOME", "NOME DO ALUNO", "ALUNO",
    "TOTAL", "TOTAIS", "LEGENDA", "OBSERVACAO", "OBSERVACOES",
}


def _text(value):
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def _normalized_text(value):
    text = unicodedata.normalize("NFKD", _text(value))
    return "".join(char for char in text if not unicodedata.combining(char)).upper()


def _header_key(value):
    return re.sub(r"[^A-Z0-9]", "", _normalized_text(value))


def _class_sheet_layout(worksheet):
    rows = list(worksheet.iter_rows(min_row=6, max_row=8, values_only=True))
    if len(rows) != 3 or len(rows[0]) < 3:
        return None
    turma = _text(rows[0][2])
    title = re.sub(r"[°()]", " ", _normalized_text(turma))
    # O ordinal (3º), ANO e os sufixos de projeto/integral sao opcionais.
    if not re.fullmatch(r"\s*\d{1,2}\s*O?\s*(?:ANO\s*)?[A-Z](?:\s*[PI])?\s*", title):
        return None

    columns = {}
    for index, value in enumerate(rows[2]):
        key = _header_key(value)
        for field, aliases in _HEADER_ALIASES.items():
            if key in aliases and field not in columns:
                columns[field] = index
    if not _REQUIRED_HEADERS.issubset(columns):
        return None
    return turma, columns


def _row_value(row, columns, field):
    index = columns.get(field)
    return row[index] if index is not None and index < len(row) else None


def _is_student_row(row, columns):
    # Numeracao e nome descartam linhas reservadas, totais e legendas do modelo.
    number = _text(_row_value(row, columns, "numero"))
    if not re.fullmatch(r"\d+(?:\.0+)?", number) or float(number) <= 0:
        return False
    name = _normalized_text(_row_value(row, columns, "nome"))
    return bool(
        name
        and name not in _INVALID_NAMES
        and not name.startswith(("#", "=", "TOTAL ", "TOTAIS ", "LEGENDA "))
    )


def read_class_sheet_records(file_obj):
    """Retorna registros brutos das abas visiveis ou None para o leitor legado.

    O layout e identificado pelo titulo em C6 e pelos cabecalhos na linha 8.
    Uma lista vazia indica layout reconhecido, mas sem alunos em abas visiveis;
    isso impede que o fallback importe por engano a primeira aba oculta.
    Streams permanecem abertos e com a mesma posicao para o leitor consolidado.
    Arquivos .xls seguem para o leitor legado, que utiliza xlrd.
    """
    position = file_obj.tell() if hasattr(file_obj, "tell") else None
    workbook = None
    try:
        if position is not None:
            file_obj.seek(0)
        try:
            workbook = load_workbook(file_obj, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile):
            return None

        records = []
        recognized = False
        for worksheet in workbook.worksheets:
            layout = _class_sheet_layout(worksheet)
            if layout is None:
                continue
            recognized = True
            if worksheet.sheet_state != "visible":
                continue
            turma, columns = layout
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=9, values_only=True), start=9
            ):
                if not _is_student_row(row, columns):
                    continue
                records.append({
                    "turma": turma,
                    "nome": _row_value(row, columns, "nome"),
                    "data_nascimento": _row_value(row, columns, "data_nascimento"),
                    "ra": _row_value(row, columns, "ra"),
                    "digito_ra": _row_value(row, columns, "digito_ra"),
                    "situacao": _row_value(row, columns, "situacao"),
                    "observacoes": _row_value(row, columns, "observacoes"),
                    "row_number": row_number,
                    "sheet_name": worksheet.title,
                })
        return records if recognized else None
    finally:
        if workbook is not None:
            workbook.close()
        if position is not None:
            file_obj.seek(position)
