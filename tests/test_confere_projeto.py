"""Regressoes de turmas de projeto com documentos ficticios em memoria."""

import io
import unittest
from pathlib import Path

from flask import Flask
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from confere import _build_result_groups, confere_bp
from services.confere_escolas import get_confere_school_config
from test_confere_abas import as_stream, make_class_sheet
from services.conferir_listas import (
    build_excel_report,
    compare_lista_piloto_sed,
    extract_sed_pdf_records,
    normalize_name_for_match,
    normalize_text,
    normalize_turma_lista,
    normalize_turma_sed,
    prepare_result_for_view,
    preview_sed_pdf_scope,
    ra_keys,
    read_lista_piloto,
    run_conferencia,
    status_compativel,
)


ROOT = Path(__file__).resolve().parents[1]
PDF_INFO = {"total_files": 1, "successful_files": ["sed_ficticio.pdf"], "errors": [], "duplicate_files": []}


def make_record(source, turma, identity="3"):
    name = f"ALUNO FICTICIO PROJETO {identity}"
    ra = f"99000{identity}"
    normalizer = normalize_turma_lista if source == "lista" else normalize_turma_sed
    return {
        "source": source,
        "turma": turma,
        "turma_key": normalizer(turma),
        "nome": name,
        "nome_norm": normalize_text(name),
        "nome_match_norm": normalize_name_for_match(name),
        "ra": ra,
        "ra_keys": sorted(ra_keys(ra)),
        "data_nascimento": "01/01/2017",
        "data_nascimento_norm": "01/01/2017",
        "situacao": "MA" if source == "lista" else "ATIVO",
        "observacoes": "",
        "pdf_origem": "sed_ficticio.pdf" if source == "sed" else "",
    }


def make_project_pdf(years=(3, 4)):
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=18, rightMargin=18)
    styles = getSampleStyleSheet()
    story = []
    for index, year in enumerate(years):
        if index:
            story.append(PageBreak())
        table = Table(
            [
                ["Serie", "N", "Nome do Aluno", "RA", "Digito", "UF", "Nascimento", "Movimentacao", "Situacao"],
                [str(year), "1", f"ALUNO FICTICIO PROJETO {year}", f"99000{year}", "", "SP", "01/01/2017", "", "ATIVO"],
            ],
            colWidths=[34, 24, 210, 66, 40, 30, 85, 85, 66],
        )
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ]))
        story.extend([
            Paragraph(f"Turma: {year}\u00b0 ANO AP TARDE ANUAL", styles["Normal"]),
            Spacer(1, 8),
            table,
        ])
    doc.build(story)
    return output.getvalue()


def make_project_workbook():
    output = io.BytesIO()
    wb = Workbook()
    wb.active.title = "RESUMO FICTICIO"
    ws = wb.create_sheet("Verificação SED")
    ws.append(["TURMA", "NOME DO ALUNO", "NASC.", "R.A.", "CÓD.", "OBSERVAÇÃO"])
    for year in (3, 4):
        ws.append([
            f"{year}\u00baA", f"ALUNO FICTICIO PROJETO {year}",
            "01/01/2017", f"99000{year}", "Matricula Ativa", "OBSERVACAO FICTICIA",
        ])
    wb.create_sheet("ABA NAO UTILIZADA").append(["NAO LER ESTA ABA"])
    wb.save(output)
    output.seek(0)
    return output


def make_status_workbook(statuses, class_layout=False):
    workbook = Workbook()
    worksheet = workbook.active
    if class_layout:
        worksheet.title = "RESUMO FICTICIO"
        worksheet = make_class_sheet(workbook, "3AP", "3º ANO AP")
    else:
        worksheet.title = "Verificação SED"
        worksheet.append(["TURMA", "NOME DO ALUNO", "NASC.", "R.A.", "CÓD."])
    for index, status in enumerate(statuses, start=1):
        row_number = index + (8 if class_layout else 1)
        if class_layout:
            values = {
                1: index, 2: f"9001{index:03d}", 3: f"ALUNO FICTICIO LEGENDA {index}",
                7: "01/01/2017", 8: f"9901{index:03d}", 10: status,
            }
        else:
            values = {
                1: "3ºA", 2: f"ALUNO FICTICIO LEGENDA {index}",
                3: "01/01/2017", 4: f"9901{index:03d}", 5: status,
            }
        for column, value in values.items():
            worksheet.cell(row_number, column, value)
    return as_stream(workbook)


class ConfereProjetoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project_pdf = make_project_pdf()

    def test_normalizes_project_codes_and_sed_descriptions(self):
        cases = {
            "3\u00baA": "3A",
            "3\u00baAP": "3A",
            "3\u00ba AP": "3A",
            "3\u00b0AP": "3A",
            "3AP": "3A",
            "3 AP": "3A",
            "3\u00b0 ANO AP TARDE ANUAL": "3A",
            "3\u00ba ANO AP TARDE ANUAL": "3A",
            "4AP": "4A",
            "4\u00ba AP": "4A",
            "4\u00b0 ANO AP MANHA ANUAL": "4A",
            "2\u00b0 ANO AI TARDE ANUAL": "2A",
            "2\u00b0 ANO G TARDE ANUAL": "2G",
        }
        for normalizer in (normalize_turma_lista, normalize_turma_sed):
            for text, expected in cases.items():
                with self.subTest(normalizer=normalizer.__name__, text=text):
                    self.assertEqual(normalizer(text), expected)

    def test_preserves_class_letters_that_are_also_suffixes_or_ordinals(self):
        for letter in ("P", "O", "I"):
            for text in (f"3\u00ba{letter}", f"3\u00ba {letter}", f"3{letter}", f"3\u00b0 ANO {letter} TARDE ANUAL", f"3\u00b0 ANO {letter}P TARDE ANUAL", f"3\u00b0 ANO {letter}I TARDE ANUAL"):
                for normalizer in (normalize_turma_lista, normalize_turma_sed):
                    with self.subTest(normalizer=normalizer.__name__, text=text):
                        self.assertEqual(normalizer(text), f"3{letter}")

    def test_rejects_unknown_suffixes_instead_of_merging_other_classes(self):
        for text in ("3\u00ba AX", "3AX", "3\u00b0 ANO AX TARDE ANUAL", "3\u00b0 ANO APO TARDE ANUAL", "SEM TURMA", ""):
            for normalizer in (normalize_turma_lista, normalize_turma_sed):
                with self.subTest(normalizer=normalizer.__name__, text=text):
                    self.assertEqual(normalizer(text), "")

    def test_projects_match_regular_lista_classes_without_false_alerts(self):
        result = compare_lista_piloto_sed(
            [make_record("lista", f"{year}\u00baA", str(year)) for year in (3, 4)],
            [make_record("sed", f"{year}\u00b0 ANO AP TARDE ANUAL", str(year)) for year in (3, 4)],
            PDF_INFO,
        )
        self.assertEqual(result["summary"]["total_ok"], 2)
        self.assertEqual(result["summary"]["total_divergencias_cadastrais"], 0)
        self.assertEqual(result["summary"]["total_lista_piloto"], 2)
        self.assertEqual(result["summary"]["turmas_conferidas"], "3A, 4A")
        self.assertTrue(all(not row["campos_divergentes"] for row in result["rows"]))

    def test_project_scope_includes_missing_students_only_from_uploaded_class(self):
        result = compare_lista_piloto_sed(
            [
                make_record("lista", "3\u00baA"),
                make_record("lista", "3\u00baA", "31"),
                make_record("lista", "4\u00baA", "4"),
            ],
            [make_record("sed", "3\u00b0 ANO AP TARDE ANUAL")],
            PDF_INFO,
        )
        summary = result["summary"]
        self.assertEqual(summary["turmas_conferidas"], "3A")
        self.assertEqual(summary["total_lista_piloto_geral"], 3)
        self.assertEqual(summary["total_lista_piloto"], 2)
        self.assertEqual(summary["total_ok"], 1)
        self.assertEqual(summary["total_nao_encontrados_sed"], 1)
        self.assertFalse(any(row["turma_lista"] == "4\u00baA" for row in result["rows"]))

    def test_real_class_letter_and_year_mismatches_remain_visible(self):
        for turma_sed in ("3\u00b0 ANO BP TARDE ANUAL", "4\u00b0 ANO AP TARDE ANUAL"):
            with self.subTest(turma_sed=turma_sed):
                result = compare_lista_piloto_sed(
                    [make_record("lista", "3\u00baA")],
                    [make_record("sed", turma_sed)],
                    PDF_INFO,
                )
                self.assertEqual(result["summary"]["total_ok"], 0)
                self.assertEqual(result["summary"]["total_divergencias_cadastrais"], 1)
                self.assertIn("turma", result["rows"][0]["campos_divergentes"])
                self.assertIn("turma diferente", result["rows"][0]["observacao"])

    def test_extracts_project_pdf_and_previews_both_classes(self):
        records = extract_sed_pdf_records(self.project_pdf, "sed_ficticio.pdf")
        self.assertEqual([record["turma_key"] for record in records], ["3A", "4A"])
        self.assertEqual([record["turma"] for record in records], ["3\u00b0 ANO AP TARDE ANUAL", "4\u00b0 ANO AP TARDE ANUAL"])
        preview = preview_sed_pdf_scope([{"filename": "sed_ficticio.pdf", "content": self.project_pdf}])
        self.assertEqual(preview["turmas"], ["3\u00baA", "4\u00baA"])
        self.assertEqual(preview["pdfs_validos"], 1)
        self.assertEqual(preview["errors"], [])

    def test_sed_only_result_groups_and_excel_use_class_without_project_suffix(self):
        raw_turma = "3\u00b0 ANO AP TARDE ANUAL"
        result = compare_lista_piloto_sed([], [make_record("sed", raw_turma)], PDF_INFO)
        view = prepare_result_for_view(result)
        self.assertEqual(view["rows"][0]["turma_sed_display"], "3\u00baA")
        self.assertEqual(view["rows"][0]["turma_sed"], raw_turma)
        groups = _build_result_groups(view)
        self.assertEqual([(group["key"], group["label"]) for group in groups], [("3A", "3\u00baA")])
        wb = load_workbook(build_excel_report(result))
        self.assertEqual(wb["Sem Lista"].cell(2, 1).value, "3\u00baA")
        self.assertEqual(wb["Para imprimir"].cell(2, 1).value, "3\u00baA")
        headers = [cell.value for cell in wb["Base completa"][1]]
        self.assertEqual(wb["Base completa"].cell(2, headers.index("turma_sed") + 1).value, raw_turma)

    def test_maria_nilza_and_mahatma_read_same_workbook_model_and_compare_projects(self):
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            with self.subTest(school_id=school_id):
                school = get_confere_school_config(school_id)
                self.assertIsNotNone(school)
                records = read_lista_piloto(make_project_workbook(), school_config=school)
                self.assertEqual([record["row_number"] for record in records], [2, 3])
                self.assertEqual([record["turma_key"] for record in records], ["3A", "4A"])
                self.assertEqual([record["ra"] for record in records], ["990.003", "990.004"])
                self.assertTrue(all(record["school_id"] == school_id for record in records))
                self.assertTrue(all(record["observacoes"] == "OBSERVACAO FICTICIA" for record in records))
                result = run_conferencia(
                    make_project_workbook(),
                    [{"filename": "sed_ficticio.pdf", "content": self.project_pdf}],
                    school_config=school,
                )
                self.assertEqual(result["school"]["id"], school_id)
                self.assertEqual(result["summary"]["total_ok"], 2)
                self.assertEqual(result["summary"]["total_divergencias_cadastrais"], 0)

    def test_flask_select_lists_maria_nilza_and_preview_uses_project_classes(self):
        app = Flask(__name__, template_folder=str(ROOT / "templates"), static_folder=str(ROOT / "static"))
        app.config.update(TESTING=True, SECRET_KEY="fictitious-test-session-key")
        app.register_blueprint(confere_bp, url_prefix="/confere")
        with app.test_client() as client:
            with client.session_transaction() as session:
                session["confere_logged_in"] = True
                session["confere_school_id"] = "maria_nilza"
            response = client.get("/confere/")
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertRegex(html, r'<option value="maria_nilza"[^>]*selected')
            self.assertIn(get_confere_school_config("maria_nilza").nome, html)
            self.assertIn('<option value="mahatma_gandhi"', html)
            response = client.post("/confere/preview-pdfs", data={
                "sed_pdfs": (io.BytesIO(self.project_pdf), "sed_ficticio.pdf"),
            })
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()["turmas"], ["3\u00baA", "4\u00baA"])

    def test_consolidated_verificacao_sheet_takes_precedence_over_class_tabs(self):
        workbook = load_workbook(make_project_workbook())
        worksheet = make_class_sheet(workbook, "5AP", "5º ANO AP")
        worksheet["C9"] = "ALUNO FICTICIO DA ABA AUXILIAR"
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            with self.subTest(school_id=school_id):
                records = read_lista_piloto(
                    as_stream(workbook), school_config=get_confere_school_config(school_id),
                )
                self.assertEqual([record["turma_key"] for record in records], ["3A", "4A"])
                self.assertEqual([record["row_number"] for record in records], [2, 3])

    def test_maria_nilza_prefers_its_conferencia_sed_sheet(self):
        workbook = load_workbook(make_project_workbook())
        worksheet = workbook.create_sheet("Conferência SED")
        worksheet.append(["Nº", "RM", "NOME DO ALUNO", "", "", "NASC.", "R.A.", "CÓD.", "OBSERVAÇÃO", "TURMA"])
        worksheet.append([1, "900001", "ALUNO FICTICIO PROJETO 3", "", "", "01/01/2017", "990003", "M.N.", "", "3ºA"])
        make_class_sheet(workbook, "5AP", "5º ANO AP")["C9"] = "ALUNO FICTICIO DA ABA AUXILIAR"

        for sheet_title in ("Conferência SED", "Conferencia sed"):
            with self.subTest(sheet_title=sheet_title):
                worksheet.title = sheet_title
                records = read_lista_piloto(
                    as_stream(workbook), school_config=get_confere_school_config("maria_nilza"),
                )
                self.assertEqual(len(records), 1)
                self.assertEqual(records[0]["turma_key"], "3A")
                self.assertEqual(records[0]["situacao"], "MA")
                self.assertEqual(records[0]["sheet_name"], "")
                result = run_conferencia(
                    as_stream(workbook),
                    [{"filename": "sed_ficticio.pdf", "content": self.project_pdf}],
                    school_config=get_confere_school_config("maria_nilza"),
                )
                self.assertEqual(result["summary"]["total_ok"], 1)
                self.assertEqual(result["summary"]["total_lista_piloto_geral"], 1)
                self.assertEqual(result["summary"]["total_divergencias_cadastrais"], 0)

        mahatma_records = read_lista_piloto(
            as_stream(workbook), school_config=get_confere_school_config("mahatma_gandhi"),
        )
        self.assertEqual([record["turma_key"] for record in mahatma_records], ["3A", "4A"])

    def test_class_sheets_mn_status_matches_active_sed_for_both_schools(self):
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            for status in ("M.N.", "MN", "M. N."):
                with self.subTest(school_id=school_id, status=status):
                    workbook = Workbook()
                    make_class_sheet(workbook, "1AI", "1º ANO AI", state="hidden")
                    for year in (3, 4):
                        sheet = make_class_sheet(workbook, f"{year}A(P)", f"{year}º ANO A(P)")
                        sheet["C9"] = f"ALUNO FICTICIO PROJETO {year}"
                        sheet["G9"] = "01/01/2017"
                        sheet["H9"] = f"99000{year}"
                        sheet["J9"] = status
                    result = run_conferencia(
                        as_stream(workbook),
                        [{"filename": "sed_ficticio.pdf", "content": self.project_pdf}],
                        school_config=get_confere_school_config(school_id),
                    )
                    self.assertEqual(result["summary"]["total_ok"], 2)
                    self.assertEqual(result["summary"]["total_lista_piloto_geral"], 2)
                    self.assertEqual(result["summary"]["total_inconsistencias_situacao"], 0)
                    self.assertTrue(all(row["situacao_lista"] == "MA" for row in result["rows"]))

    def test_school_legend_statuses_are_normalized_in_both_workbook_layouts(self):
        cases = [
            ("TE", "TE"), ("T.E.", "TE"), ("T. E.", "TE"),
            ("TR", "MA"), ("T.R.", "MA"), ("T. R.", "MA"),
            ("Transferência Expedida", "TE"),
            ("MN", "MA"), ("M.N.", "MA"), ("M. N.", "MA"),
            ("Matrícula Normal", "MA"),
            ("REM", "REM"), ("Remanejado", "REM"),
            ("NF", "NF"), ("N.F.", "NF"), ("Não Frequente", "NF"),
            ("NCOM", "NCOM"), ("Não Compareceu", "NCOM"),
            ("", "MA"), (0, "MA"), ("PNEE", "MA"),
        ]
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            for class_layout in (False, True):
                with self.subTest(school_id=school_id, class_layout=class_layout):
                    records = read_lista_piloto(
                        make_status_workbook([status for status, _ in cases], class_layout),
                        school_config=get_confere_school_config(school_id),
                    )
                    self.assertEqual([record["situacao"] for record in records], [expected for _, expected in cases])

    def test_transfer_legend_distinguishes_received_from_expedited(self):
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            for class_layout in (False, True):
                records = read_lista_piloto(
                    make_status_workbook(["TE", "T.E.", "TR", "T.R."], class_layout=class_layout),
                    school_config=get_confere_school_config(school_id),
                )
                self.assertEqual([record["situacao"] for record in records], ["TE", "TE", "MA", "MA"])
                for sed_status, expected_ok in (("BXTR", 2), ("TRANSF", 2), ("ATIVO", 2)):
                    with self.subTest(school_id=school_id, class_layout=class_layout, sed_status=sed_status):
                        sed_records = [
                            {**record, "source": "sed", "situacao": sed_status, "pdf_origem": "sed_ficticio.pdf"}
                            for record in records
                        ]
                        result = compare_lista_piloto_sed(records, sed_records, PDF_INFO)
                        self.assertEqual(result["summary"]["total_ok"], expected_ok)
                        self.assertEqual(result["summary"]["total_inconsistencias_situacao"], 4 - expected_ok)

    def test_attendance_codes_remain_pending_with_readable_view_and_excel_labels(self):
        labels = {"NF": "Não frequente", "NCOM": "Não compareceu"}
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            with self.subTest(school_id=school_id):
                records = read_lista_piloto(
                    make_status_workbook(["NF", "NCOM"], class_layout=True),
                    school_config=get_confere_school_config(school_id),
                )
                for record in records:
                    self.assertFalse(status_compativel(record["situacao"], "ATIVO"))
                    self.assertFalse(status_compativel(record["situacao"], "BXTR"))
                sed_records = [
                    {**record, "source": "sed", "situacao": "ATIVO", "pdf_origem": "sed_ficticio.pdf"}
                    for record in records
                ]
                result = prepare_result_for_view(compare_lista_piloto_sed(records, sed_records, PDF_INFO))
                self.assertEqual(result["summary"]["total_ok"], 0)
                self.assertEqual(result["summary"]["total_inconsistencias_situacao"], 2)
                for row in result["rows"]:
                    self.assertIn("situacao", row["campos_divergentes"])
                    label = labels[row["situacao_lista"]]
                    self.assertEqual(normalize_text(row["situacao_lista_display"]), normalize_text(label))
                    self.assertIn(normalize_text(label), normalize_text(row["observacao"]))
                workbook = load_workbook(build_excel_report(result))
                worksheet = workbook["Inconsistencias"]
                headers = [cell.value for cell in worksheet[1]]
                column = headers.index("Sit. Lista") + 1
                excel_statuses = {
                    normalize_text(worksheet.cell(row, column).value)
                    for row in range(2, worksheet.max_row + 1)
                }
                self.assertEqual(excel_statuses, {"NAO FREQUENTE", "NAO COMPARECEU"})

    def test_no_show_status_matches_ncom_sed_and_non_frequent_stays_pending(self):
        self.assertTrue(status_compativel("NCOM", "Não Comparecimento"))
        for school_id in ("maria_nilza", "mahatma_gandhi"):
            with self.subTest(school_id=school_id):
                records = read_lista_piloto(
                    make_status_workbook(["NCOM", "NF"], class_layout=True),
                    school_config=get_confere_school_config(school_id),
                )
                sed_records = [
                    {**record, "source": "sed", "situacao": "NCOM", "pdf_origem": "sed_ficticio.pdf"}
                    for record in records
                ]
                result = compare_lista_piloto_sed(records, sed_records, PDF_INFO)
                self.assertEqual(result["summary"]["total_ok"], 1)
                self.assertEqual(result["summary"]["total_inconsistencias_situacao"], 1)
                self.assertEqual(next(row for row in result["rows"] if row["categoria"] == "ok")["situacao_lista"], "NCOM")

    def test_school_specific_transfer_and_enrollment_aliases_do_not_change_padin(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "LISTA CORRIDA"
        worksheet.append(["TURMA", "NOME", "DATA NASC", "RA", "COD"])
        worksheet.append(["3ºA", "ALUNO FICTICIO TR", "01/01/2017", "9901001", "TR"])
        worksheet.append(["3ºA", "ALUNO FICTICIO MN", "01/01/2017", "9901002", "MN"])
        records = read_lista_piloto(as_stream(workbook), school_config=get_confere_school_config("padin"))
        self.assertEqual([record["situacao"] for record in records], ["TR", "MN"])


if __name__ == "__main__":
    unittest.main()
