import io
import unittest
from datetime import datetime

from openpyxl import Workbook

from services.confere_abas import read_class_sheet_records


def make_class_sheet(workbook, title, turma, status_column=10, state="visible"):
    sheet = workbook.create_sheet(title)
    sheet.sheet_state = state
    sheet["C6"] = turma
    for column, header in {1: "Nº", 2: "RM", 3: "NOME DO ALUNO", 4: "SEXO", 7: "NASC.", 8: "R.A."}.items():
        sheet.cell(8, column, header)
    sheet.cell(8, status_column, "CÓD.")
    sheet.cell(8, status_column + 1, "OBSERVAÇÃO")
    for column, value in {
        1: 1, 2: "900001", 3: "ALUNO FICTICIO TESTE", 4: "F",
        7: datetime(2017, 1, 2), 8: "990000001", status_column: "MA",
        status_column + 1: "EXEMPLO FICTICIO",
    }.items():
        sheet.cell(9, column, value)
    return sheet


def as_stream(workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


class ClassSheetReaderTests(unittest.TestCase):
    def test_reads_visible_classes_and_ignores_hidden_and_auxiliary_sheets(self):
        workbook = Workbook()
        workbook.active.title = "RESUMO"
        hidden = make_class_sheet(workbook, "1AI", "1º ANO AI", status_column=9, state="hidden")
        hidden["C9"] = "ALUNO FICTICIO OCULTO"
        make_class_sheet(workbook, "3A(P)", "3º ANO A(P)")
        second = make_class_sheet(workbook, "4A(P)", "4º ANO AP")
        second["C9"] = "OUTRO ALUNO FICTICIO"
        make_class_sheet(workbook, "2AI", "2º ANO AI", state="veryHidden")
        records = read_class_sheet_records(as_stream(workbook))
        self.assertEqual([record["turma"] for record in records], ["3º ANO A(P)", "4º ANO AP"])
        self.assertEqual([record["sheet_name"] for record in records], ["3A(P)", "4A(P)"])
        self.assertEqual([record["row_number"] for record in records], [9, 9])
        self.assertEqual(records[0]["data_nascimento"], datetime(2017, 1, 2))
        self.assertEqual(records[0]["ra"], "990000001")
        self.assertEqual(records[0]["situacao"], "MA")
        self.assertEqual(records[0]["observacoes"], "EXEMPLO FICTICIO")

    def test_matches_status_and_observation_headers_when_columns_shift(self):
        workbook = Workbook()
        sheet = make_class_sheet(workbook, "1AI", "1º ANO AI", status_column=9)
        sheet["I9"] = "TE"
        sheet["J9"] = "TRANSFERENCIA FICTICIA"
        records = read_class_sheet_records(as_stream(workbook))
        self.assertEqual(records[0]["situacao"], "TE")
        self.assertEqual(records[0]["observacoes"], "TRANSFERENCIA FICTICIA")
        self.assertIsNone(records[0]["digito_ra"])

    def test_uses_ra_digit_only_with_an_explicit_header(self):
        workbook = Workbook()
        sheet = make_class_sheet(workbook, "3AP", "3º ANO AP")
        sheet["I9"] = "X"
        self.assertIsNone(read_class_sheet_records(as_stream(workbook))[0]["digito_ra"])
        sheet["I8"] = "DÍGITO RA"
        self.assertEqual(read_class_sheet_records(as_stream(workbook))[0]["digito_ra"], "X")

    def test_skips_reserved_rows_totals_legends_and_formula_errors(self):
        workbook = Workbook()
        sheet = make_class_sheet(workbook, "3AP", "3º ANO A(P)")
        for row_number, name in enumerate([None, "0", "#REF!", "#REF#", "nan", "TOTAL DE ALUNOS", "LEGENDA", "NOME DO ALUNO"], start=10):
            sheet.cell(row_number, 1, row_number - 8)
            sheet.cell(row_number, 3, name)
        sheet["C19"] = "LEGENDA SEM NUMERACAO"
        sheet["A20"] = "TOTAL"
        sheet["C20"] = "ALUNOS MATRICULADOS"
        sheet["A21"] = 0
        sheet["C21"] = "OUTRO TOTAL"
        sheet["A22"] = 2
        sheet["C22"] = '=""'
        records = read_class_sheet_records(as_stream(workbook))
        self.assertEqual(len(records), 1)

    def test_returns_none_for_legacy_layout_and_preserves_stream_position(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A8"] = "ALUNO FICTICIO LEGADO"
        sheet["D8"] = "3ºA"
        stream = as_stream(workbook)
        stream.seek(12)
        self.assertIsNone(read_class_sheet_records(stream))
        self.assertEqual(stream.tell(), 12)
        self.assertFalse(stream.closed)

    def test_returns_empty_when_only_hidden_sheets_have_the_known_layout(self):
        workbook = Workbook()
        make_class_sheet(workbook, "3AP", "3º ANO AP", state="hidden")
        self.assertEqual(read_class_sheet_records(as_stream(workbook)), [])

    def test_requires_class_title_and_known_headers(self):
        workbook = Workbook()
        sheet = make_class_sheet(workbook, "RESUMO", "TOTAL DE ALUNOS")
        self.assertIsNone(read_class_sheet_records(as_stream(workbook)))
        sheet["C6"] = "3º ANO AP"
        sheet["H8"] = "OUTRA COLUNA"
        self.assertIsNone(read_class_sheet_records(as_stream(workbook)))

    def test_leaves_legacy_xls_stream_to_fallback_without_consuming_it(self):
        stream = io.BytesIO(bytes.fromhex("D0CF11E0A1B11AE1") + b"synthetic xls fixture")
        stream.seek(4)
        self.assertIsNone(read_class_sheet_records(stream))
        self.assertEqual(stream.tell(), 4)

    def test_reads_xlsm_named_stream_without_modifying_it(self):
        workbook = Workbook()
        make_class_sheet(workbook, "3AP", "3º ANO AP")
        stream = as_stream(workbook)
        stream.name = "lista_ficticia.xlsm"
        original = stream.getvalue()
        self.assertEqual(len(read_class_sheet_records(stream)), 1)
        self.assertEqual(stream.getvalue(), original)
        self.assertEqual(stream.tell(), 0)


if __name__ == "__main__":
    unittest.main()
